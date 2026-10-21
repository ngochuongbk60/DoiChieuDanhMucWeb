# -*- coding: utf-8 -*-
"""
PHẦN MỀM ĐỐI CHIẾU DANH MỤC KỸ THUẬT VÀ TỔNG HỢP KẾT QUẢ
Phiên bản WEB APP (Flask)

File gốc: Phụ lục XII - Điều dưỡng (TT32/TT-BYT)
Đối chiếu với 05 file còn lại.
"""

import os
import re
import json
import uuid
import time
import unicodedata
import threading
from datetime import datetime
from difflib import SequenceMatcher

from flask import Flask, request, jsonify, send_file, Response, send_from_directory
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

app = Flask(__name__, static_folder='static')
app.config['MAX_CONTENT_LENGTH'] = 100 * 1024 * 1024  # 100MB max

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), 'uploads')
RESULT_DIR = os.path.join(os.path.dirname(__file__), 'results')
os.makedirs(UPLOAD_DIR, exist_ok=True)
os.makedirs(RESULT_DIR, exist_ok=True)

# Store active sessions' progress
sessions = {}


# ============================================================================
# CÁC HÀM TIỆN ÍCH
# ============================================================================

def normalize_text(text):
    """Chuẩn hóa văn bản để so sánh."""
    if text is None:
        return ""
    text = str(text).strip()
    text = text.replace('\n', ' ').replace('\r', ' ')
    text = re.sub(r'\s+', ' ', text)
    text = unicodedata.normalize('NFC', text)
    text = text.lower()
    text = text.rstrip('*+ ')
    return text


def safe_str(value):
    """Chuyển đổi giá trị về chuỗi an toàn."""
    if value is None:
        return ""
    return str(value).strip()


def read_excel_data(filepath, sheet_name=None):
    """Đọc workbook từ file Excel."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = wb.active
    return wb, ws


# ============================================================================
# CHIẾN LƯỢC SO KHỚP TỐI ƯU
# ============================================================================

EXACT_MATCH_THRESHOLD = 0.92
PARTIAL_MATCH_THRESHOLD = 0.50


def build_lookup(records, name_field_normalized):
    """Xây dựng lookup dict từ danh sách records."""
    exact_dict = {}
    all_records = []
    for rec in records:
        norm_name = rec.get(name_field_normalized, '')
        if not norm_name:
            continue
        if norm_name not in exact_dict:
            exact_dict[norm_name] = []
        exact_dict[norm_name].append(rec)
        all_records.append(rec)
    return exact_dict, all_records


def find_matches_optimized(ten_goc_norm, exact_dict, all_records, name_field_normalized):
    """Tìm bản ghi trùng khớp."""
    if not ten_goc_norm:
        return []

    if ten_goc_norm in exact_dict:
        return [(rec, 'exact', 1.0) for rec in exact_dict[ten_goc_norm]]

    words_goc = set(w for w in ten_goc_norm.split() if len(w) >= 3)
    if not words_goc:
        words_goc = set(ten_goc_norm.split())

    candidates = []
    for rec in all_records:
        rec_name = rec.get(name_field_normalized, '')
        if not rec_name:
            continue
        words_rec = set(w for w in rec_name.split() if len(w) >= 3)
        if not words_rec:
            words_rec = set(rec_name.split())
        if words_goc and words_rec:
            common = words_goc & words_rec
            overlap = len(common) / max(len(words_goc), 1)
            if overlap >= 0.3:
                candidates.append(rec)

    exact_matches = []
    partial_matches = []
    for rec in candidates:
        rec_name = rec.get(name_field_normalized, '')
        ratio = SequenceMatcher(None, ten_goc_norm, rec_name).ratio()
        if ratio >= EXACT_MATCH_THRESHOLD:
            exact_matches.append((rec, 'exact', ratio))
        elif ratio >= PARTIAL_MATCH_THRESHOLD:
            partial_matches.append((rec, 'partial', ratio))

    if exact_matches:
        exact_matches.sort(key=lambda x: x[2], reverse=True)
        return exact_matches
    if partial_matches:
        partial_matches.sort(key=lambda x: x[2], reverse=True)
        return partial_matches[:5]
    return []


# ============================================================================
# ĐỌC DỮ LIỆU
# ============================================================================

def read_file_goc(filepath):
    _, ws = read_excel_data(filepath)
    records = []
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=False):
        a_val = safe_str(row[0].value)
        b_val = safe_str(row[1].value)
        c_val = safe_str(row[2].value)
        d_val = safe_str(row[3].value)
        e_val = safe_str(row[4].value)
        if c_val and a_val:
            try:
                int(a_val)
            except ValueError:
                continue
            records.append({
                'stt_goc': a_val,
                'ma_tt43_21': b_val,
                'ten_ky_thuat_goc': c_val,
                'ten_ky_thuat_goc_norm': normalize_text(c_val),
                'bs_chi_dinh': d_val,
                'dd_chi_dinh': e_val,
            })
    return records


def read_file_1(filepath):
    _, ws = read_excel_data(filepath)
    records = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=False):
        a_val = row[0].value
        b_val = safe_str(row[1].value)
        d_val = safe_str(row[3].value)
        if d_val and a_val is not None:
            records.append({
                'ma_ky_thuat': b_val,
                'ten_ky_thuat': d_val,
                'norm': normalize_text(d_val),
            })
    return records


def read_file_2(filepath):
    _, ws = read_excel_data(filepath)
    records = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=False):
        a_val = row[0].value
        d_val = safe_str(row[3].value)
        e_val = safe_str(row[4].value)
        if e_val and a_val is not None:
            records.append({
                'ma_lien_ket': d_val,
                'ten_ky_thuat': e_val,
                'norm': normalize_text(e_val),
            })
    return records


def read_file_3(filepath):
    _, ws = read_excel_data(filepath)
    records = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=False):
        a_val = row[0].value
        c_val = safe_str(row[2].value)
        d_val = safe_str(row[3].value)
        e_val = safe_str(row[4].value)
        f_val = safe_str(row[5].value)
        g_val = safe_str(row[6].value)
        if e_val and a_val is not None:
            records.append({
                'ma_tuong_duong': c_val,
                'ma_tt43_50_21': d_val,
                'ten_ky_thuat': e_val,
                'norm': normalize_text(e_val),
                'phan_tuyen': f_val,
                'phan_loai_pttt': g_val,
            })
    return records


def read_file_4(filepath):
    _, ws = read_excel_data(filepath, sheet_name='PL03')
    records = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=False):
        a_val = row[0].value
        b_val = safe_str(row[1].value)
        c_val = safe_str(row[2].value)
        d_val = safe_str(row[3].value)
        e_val = safe_str(row[4].value)
        f_val = safe_str(row[5].value)
        i_val = safe_str(row[8].value)
        if d_val and a_val is not None:
            records.append({
                'stt': safe_str(a_val),
                'ma_dvkt': b_val,
                'ma_tuong_duong_dvkt': c_val,
                'ten_dvkt': d_val,
                'norm': normalize_text(d_val),
                'qd_phe_duyet': e_val,
                'ten_dvkt_gia': f_val,
                'qd_quy_trinh': i_val,
            })
    return records


def read_file_5(filepath):
    _, ws = read_excel_data(filepath)
    records = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=False):
        a_val = row[0].value
        b_val = safe_str(row[1].value)
        if b_val and a_val is not None:
            records.append({
                'stt': safe_str(a_val),
                'ten_quy_trinh': b_val,
                'norm': normalize_text(b_val),
            })
    return records


# ============================================================================
# TẠO BÁO CÁO EXCEL
# ============================================================================

def create_result_workbook(goc_records, data_f1, data_f2, data_f3, data_f4, data_f5,
                           output_path, log_callback=None, progress_callback=None):
    """Tạo workbook kết quả tổng hợp."""

    def log(msg):
        if log_callback:
            log_callback(msg)

    log("  Xây dựng lookup tables...")
    lookup_f1, all_f1 = build_lookup(data_f1, 'norm')
    lookup_f2, all_f2 = build_lookup(data_f2, 'norm')
    lookup_f3, all_f3 = build_lookup(data_f3, 'norm')
    lookup_f4, all_f4 = build_lookup(data_f4, 'norm')
    lookup_f5, all_f5 = build_lookup(data_f5, 'norm')

    wb = openpyxl.Workbook()

    # ---- STYLES ----
    header_font = Font(name='Times New Roman', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')
    header_fill_goc = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_fill_f1 = PatternFill(start_color='548235', end_color='548235', fill_type='solid')
    header_fill_f2 = PatternFill(start_color='BF8F00', end_color='BF8F00', fill_type='solid')
    header_fill_f3 = PatternFill(start_color='C55A11', end_color='C55A11', fill_type='solid')
    header_fill_f4 = PatternFill(start_color='7030A0', end_color='7030A0', fill_type='solid')
    header_fill_f5 = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
    header_fill_eval = PatternFill(start_color='404040', end_color='404040', fill_type='solid')

    data_font = Font(name='Times New Roman', size=11)
    data_font_bold = Font(name='Times New Roman', size=11, bold=True)
    wrap_alignment = Alignment(wrap_text=True, vertical='top')
    center_alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)

    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    light_green = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    light_red = PatternFill(start_color='FCE4EC', end_color='FCE4EC', fill_type='solid')
    light_orange = PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid')
    light_yellow = PatternFill(start_color='FFFDE7', end_color='FFFDE7', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    # ================================================================
    # SHEET 1: KẾT QUẢ TỔNG HỢP
    # ================================================================
    ws = wb.active
    ws.title = "KET QUA DOI CHIEU"

    ws.merge_cells('A1:AC1')
    ws['A1'] = 'BẢNG ĐỐI CHIẾU DANH MỤC KỸ THUẬT CHUYÊN MÔN - PHỤ LỤC XII (TT32/TT-BYT)'
    ws['A1'].font = Font(name='Times New Roman', bold=True, size=14, color='2F5496')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells('A2:AC2')
    ws['A2'] = f'Ngày tạo báo cáo: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}'
    ws['A2'].font = Font(name='Times New Roman', italic=True, size=10)
    ws['A2'].alignment = Alignment(horizontal='center')

    group_headers = [
        ('A4:E4', 'FILE GỐC - PHỤ LỤC XII (TT32)', header_fill_goc),
        ('F4:G4', 'PL01 TT23', header_fill_f1),
        ('H4:I4', 'PL02 TT23', header_fill_f2),
        ('J4:N4', 'DMKT TƯƠNG ĐƯƠNG (TT43/50/21)', header_fill_f3),
        ('O4:U4', 'PL03 DVKT', header_fill_f4),
        ('V4:W4', 'TỔNG HỢP QUY TRÌNH', header_fill_f5),
        ('X4:AC4', 'ĐÁNH GIÁ & CẢNH BÁO', header_fill_eval),
    ]
    for cell_range, title, fill in group_headers:
        ws.merge_cells(cell_range)
        top_left = cell_range.split(':')[0]
        ws[top_left] = title
        ws[top_left].font = header_font
        ws[top_left].fill = fill
        ws[top_left].alignment = Alignment(horizontal='center', vertical='center')

    headers = [
        ('A', 'STT\n(Gốc)', header_fill_goc),
        ('B', 'Mã TT 43, 21\n(Gốc)', header_fill_goc),
        ('C', 'Danh mục kỹ thuật\n(Gốc)', header_fill_goc),
        ('D', 'BS chỉ định\n(Gốc)', header_fill_goc),
        ('E', 'ĐD chỉ định\n(Gốc)', header_fill_goc),
        ('F', 'Mã kỹ thuật\n(PL01-TT23)', header_fill_f1),
        ('G', 'Tên kỹ thuật\n(PL01-TT23)', header_fill_f1),
        ('H', 'Mã liên kết\n(PL02-TT23)', header_fill_f2),
        ('I', 'Tên kỹ thuật\n(PL02-TT23)', header_fill_f2),
        ('J', 'Mã tương đương\n(TT43/50/21)', header_fill_f3),
        ('K', 'Mã TT43, 50, 21', header_fill_f3),
        ('L', 'Tên DMKT\n(TT43/50/21)', header_fill_f3),
        ('M', 'Phân tuyến', header_fill_f3),
        ('N', 'Phân loại PTTT', header_fill_f3),
        ('O', 'STT\n(PL03)', header_fill_f4),
        ('P', 'Mã DVKT\n(PL03)', header_fill_f4),
        ('Q', 'Mã tương đương\nDVKT (PL03)', header_fill_f4),
        ('R', 'Tên DVKT phê duyệt\n(PL03)', header_fill_f4),
        ('S', 'QĐ phê duyệt\nthực hiện (PL03)', header_fill_f4),
        ('T', 'Tên DVKT phê duyệt\ngiá (PL03)', header_fill_f4),
        ('U', 'QĐ ban hành\nquy trình CM (PL03)', header_fill_f4),
        ('V', 'STT\n(QT đã BH)', header_fill_f5),
        ('W', 'Tên quy trình\n(QT đã BH)', header_fill_f5),
        ('X', 'Trạng thái\nPL01-TT23', header_fill_eval),
        ('Y', 'Trạng thái\nPL02-TT23', header_fill_eval),
        ('Z', 'Trạng thái\nDMKT-TĐ', header_fill_eval),
        ('AA', 'Trạng thái\nPL03-DVKT', header_fill_eval),
        ('AB', 'Trạng thái\nQT-BH', header_fill_eval),
        ('AC', 'CẢNH BÁO\nTỔNG HỢP', header_fill_eval),
    ]

    for col_letter, header_text, fill in headers:
        cell = ws[f'{col_letter}5']
        cell.value = header_text
        cell.font = header_font
        cell.fill = fill
        cell.alignment = center_alignment
        cell.border = thin_border

    ws.freeze_panes = 'A6'

    # ---- THỰC HIỆN ĐỐI CHIẾU ----
    stats = {
        'total': len(goc_records),
        'f1_found': 0, 'f1_not_found': 0, 'f1_partial': 0, 'f1_multiple': 0,
        'f2_found': 0, 'f2_not_found': 0, 'f2_partial': 0, 'f2_multiple': 0,
        'f3_found': 0, 'f3_not_found': 0, 'f3_partial': 0, 'f3_multiple': 0,
        'f4_found': 0, 'f4_not_found': 0, 'f4_partial': 0, 'f4_multiple': 0,
        'f5_found': 0, 'f5_not_found': 0, 'f5_partial': 0, 'f5_multiple': 0,
        'need_review': 0, 'missing_approval': 0,
    }

    detail_rows = []
    multi_detail = []
    total = len(goc_records)

    # Also build JSON result for web display
    web_results = []

    for idx, rec in enumerate(goc_records):
        row_num = idx + 6
        ten_goc = rec['ten_ky_thuat_goc']
        ten_goc_norm = rec['ten_ky_thuat_goc_norm']
        warnings = []

        ws[f'A{row_num}'] = rec['stt_goc']
        ws[f'B{row_num}'] = rec['ma_tt43_21']
        ws[f'C{row_num}'] = rec['ten_ky_thuat_goc']
        ws[f'D{row_num}'] = rec['bs_chi_dinh']
        ws[f'E{row_num}'] = rec['dd_chi_dinh']
        ws[f'C{row_num}'].fill = yellow_fill

        web_row = {
            'stt': rec['stt_goc'],
            'ma': rec['ma_tt43_21'],
            'ten': ten_goc,
            'bs': rec['bs_chi_dinh'],
            'dd': rec['dd_chi_dinh'],
            'f1_status': '', 'f2_status': '', 'f3_status': '', 'f4_status': '', 'f5_status': '',
            'f1_type': '', 'f2_type': '', 'f3_type': '', 'f4_type': '', 'f5_type': '',
            'warnings': [],
        }

        # ---- File 1 ----
        matches = find_matches_optimized(ten_goc_norm, lookup_f1, all_f1, 'norm')
        if matches:
            best = matches[0]
            ws[f'F{row_num}'] = best[0]['ma_ky_thuat']
            ws[f'G{row_num}'] = best[0]['ten_ky_thuat']
            ws[f'F{row_num}'].font = data_font_bold
            ws[f'G{row_num}'].font = data_font_bold
            if best[1] == 'exact':
                ws[f'X{row_num}'] = '✓ Trùng khớp'
                ws[f'X{row_num}'].fill = light_green
                stats['f1_found'] += 1
                web_row['f1_status'] = '✓ Trùng khớp'
                web_row['f1_type'] = 'exact'
            else:
                pct = f"{best[2]*100:.0f}%"
                ws[f'X{row_num}'] = f'⚠ Gần khớp ({pct})'
                ws[f'X{row_num}'].fill = light_orange
                stats['f1_partial'] += 1
                warnings.append(f'PL01: Trùng khớp một phần ({pct})')
                web_row['f1_status'] = f'⚠ Gần khớp ({pct})'
                web_row['f1_type'] = 'partial'
            if len(matches) > 1:
                ws[f'X{row_num}'] = f'⚠ Nhiều KQ ({len(matches)})'
                ws[f'X{row_num}'].fill = light_orange
                stats['f1_multiple'] += 1
                warnings.append(f'PL01: {len(matches)} kết quả khớp')
                web_row['f1_status'] = f'⚠ Nhiều KQ ({len(matches)})'
                web_row['f1_type'] = 'multiple'
                for m in matches:
                    multi_detail.append((rec['stt_goc'], ten_goc, 'PL01-TT23',
                                         m[0]['ten_ky_thuat'], m[0]['ma_ky_thuat'], m[2]))
        else:
            ws[f'X{row_num}'] = '✗ Chưa tìm thấy'
            ws[f'X{row_num}'].fill = light_red
            stats['f1_not_found'] += 1
            warnings.append('PL01: Chưa tìm thấy')
            web_row['f1_status'] = '✗ Chưa tìm thấy'
            web_row['f1_type'] = 'not_found'

        # ---- File 2 ----
        matches = find_matches_optimized(ten_goc_norm, lookup_f2, all_f2, 'norm')
        if matches:
            best = matches[0]
            ws[f'H{row_num}'] = best[0]['ma_lien_ket']
            ws[f'I{row_num}'] = best[0]['ten_ky_thuat']
            ws[f'H{row_num}'].font = data_font_bold
            ws[f'I{row_num}'].font = data_font_bold
            if best[1] == 'exact':
                ws[f'Y{row_num}'] = '✓ Trùng khớp'
                ws[f'Y{row_num}'].fill = light_green
                stats['f2_found'] += 1
                web_row['f2_status'] = '✓ Trùng khớp'
                web_row['f2_type'] = 'exact'
            else:
                pct = f"{best[2]*100:.0f}%"
                ws[f'Y{row_num}'] = f'⚠ Gần khớp ({pct})'
                ws[f'Y{row_num}'].fill = light_orange
                stats['f2_partial'] += 1
                warnings.append(f'PL02: Trùng khớp một phần ({pct})')
                web_row['f2_status'] = f'⚠ Gần khớp ({pct})'
                web_row['f2_type'] = 'partial'
            if len(matches) > 1:
                ws[f'Y{row_num}'] = f'⚠ Nhiều KQ ({len(matches)})'
                ws[f'Y{row_num}'].fill = light_orange
                stats['f2_multiple'] += 1
                warnings.append(f'PL02: {len(matches)} kết quả khớp')
                web_row['f2_status'] = f'⚠ Nhiều KQ ({len(matches)})'
                web_row['f2_type'] = 'multiple'
                for m in matches:
                    multi_detail.append((rec['stt_goc'], ten_goc, 'PL02-TT23',
                                         m[0]['ten_ky_thuat'], m[0]['ma_lien_ket'], m[2]))
        else:
            ws[f'Y{row_num}'] = '✗ Chưa tìm thấy'
            ws[f'Y{row_num}'].fill = light_red
            stats['f2_not_found'] += 1
            warnings.append('PL02: Chưa tìm thấy')
            web_row['f2_status'] = '✗ Chưa tìm thấy'
            web_row['f2_type'] = 'not_found'

        # ---- File 3 ----
        matches = find_matches_optimized(ten_goc_norm, lookup_f3, all_f3, 'norm')
        if matches:
            best = matches[0]
            ws[f'J{row_num}'] = best[0]['ma_tuong_duong']
            ws[f'K{row_num}'] = best[0]['ma_tt43_50_21']
            ws[f'L{row_num}'] = best[0]['ten_ky_thuat']
            ws[f'M{row_num}'] = best[0]['phan_tuyen']
            ws[f'N{row_num}'] = best[0]['phan_loai_pttt']
            for c in ['J', 'K', 'L', 'M', 'N']:
                ws[f'{c}{row_num}'].font = data_font_bold
            if best[1] == 'exact':
                ws[f'Z{row_num}'] = '✓ Trùng khớp'
                ws[f'Z{row_num}'].fill = light_green
                stats['f3_found'] += 1
                web_row['f3_status'] = '✓ Trùng khớp'
                web_row['f3_type'] = 'exact'
            else:
                pct = f"{best[2]*100:.0f}%"
                ws[f'Z{row_num}'] = f'⚠ Gần khớp ({pct})'
                ws[f'Z{row_num}'].fill = light_orange
                stats['f3_partial'] += 1
                warnings.append(f'DMKT-TĐ: Trùng khớp một phần ({pct})')
                web_row['f3_status'] = f'⚠ Gần khớp ({pct})'
                web_row['f3_type'] = 'partial'
            if len(matches) > 1:
                ws[f'Z{row_num}'] = f'⚠ Nhiều KQ ({len(matches)})'
                ws[f'Z{row_num}'].fill = light_orange
                stats['f3_multiple'] += 1
                warnings.append(f'DMKT-TĐ: {len(matches)} kết quả khớp')
                web_row['f3_status'] = f'⚠ Nhiều KQ ({len(matches)})'
                web_row['f3_type'] = 'multiple'
                for m in matches:
                    multi_detail.append((rec['stt_goc'], ten_goc, 'DMKT-TĐ',
                                         m[0]['ten_ky_thuat'], m[0]['ma_tuong_duong'], m[2]))
        else:
            ws[f'Z{row_num}'] = '✗ Chưa tìm thấy'
            ws[f'Z{row_num}'].fill = light_red
            stats['f3_not_found'] += 1
            warnings.append('DMKT-TĐ: Chưa tìm thấy')
            web_row['f3_status'] = '✗ Chưa tìm thấy'
            web_row['f3_type'] = 'not_found'

        # ---- File 4 ----
        matches = find_matches_optimized(ten_goc_norm, lookup_f4, all_f4, 'norm')
        if matches:
            best = matches[0]
            ws[f'O{row_num}'] = best[0]['stt']
            ws[f'P{row_num}'] = best[0]['ma_dvkt']
            ws[f'Q{row_num}'] = best[0]['ma_tuong_duong_dvkt']
            ws[f'R{row_num}'] = best[0]['ten_dvkt']
            ws[f'S{row_num}'] = best[0]['qd_phe_duyet']
            ws[f'T{row_num}'] = best[0]['ten_dvkt_gia']
            ws[f'U{row_num}'] = best[0]['qd_quy_trinh']
            for c in ['O', 'P', 'Q', 'R', 'S', 'T', 'U']:
                ws[f'{c}{row_num}'].font = data_font_bold
            if best[1] == 'exact':
                ws[f'AA{row_num}'] = '✓ Trùng khớp'
                ws[f'AA{row_num}'].fill = light_green
                stats['f4_found'] += 1
                web_row['f4_status'] = '✓ Trùng khớp'
                web_row['f4_type'] = 'exact'
            else:
                pct = f"{best[2]*100:.0f}%"
                ws[f'AA{row_num}'] = f'⚠ Gần khớp ({pct})'
                ws[f'AA{row_num}'].fill = light_orange
                stats['f4_partial'] += 1
                warnings.append(f'PL03: Trùng khớp một phần ({pct})')
                web_row['f4_status'] = f'⚠ Gần khớp ({pct})'
                web_row['f4_type'] = 'partial'
            if len(matches) > 1:
                ws[f'AA{row_num}'] = f'⚠ Nhiều KQ ({len(matches)})'
                ws[f'AA{row_num}'].fill = light_orange
                stats['f4_multiple'] += 1
                warnings.append(f'PL03: {len(matches)} kết quả khớp')
                web_row['f4_status'] = f'⚠ Nhiều KQ ({len(matches)})'
                web_row['f4_type'] = 'multiple'
                for m in matches:
                    multi_detail.append((rec['stt_goc'], ten_goc, 'PL03-DVKT',
                                         m[0]['ten_dvkt'], m[0]['ma_dvkt'], m[2]))
            missing_docs = []
            if not best[0]['qd_phe_duyet']:
                missing_docs.append('QĐ phê duyệt')
            if not best[0]['ten_dvkt_gia']:
                missing_docs.append('Phê duyệt giá')
            if not best[0]['qd_quy_trinh']:
                missing_docs.append('QĐ quy trình CM')
            if missing_docs:
                warnings.append(f'PL03: Thiếu {", ".join(missing_docs)}')
                stats['missing_approval'] += 1
        else:
            ws[f'AA{row_num}'] = '✗ Chưa tìm thấy'
            ws[f'AA{row_num}'].fill = light_red
            stats['f4_not_found'] += 1
            warnings.append('PL03: Chưa tìm thấy')
            web_row['f4_status'] = '✗ Chưa tìm thấy'
            web_row['f4_type'] = 'not_found'

        # ---- File 5 ----
        matches = find_matches_optimized(ten_goc_norm, lookup_f5, all_f5, 'norm')
        if matches:
            best = matches[0]
            ws[f'V{row_num}'] = best[0]['stt']
            ws[f'W{row_num}'] = best[0]['ten_quy_trinh']
            ws[f'V{row_num}'].font = data_font_bold
            ws[f'W{row_num}'].font = data_font_bold
            if best[1] == 'exact':
                ws[f'AB{row_num}'] = '✓ Trùng khớp'
                ws[f'AB{row_num}'].fill = light_green
                stats['f5_found'] += 1
                web_row['f5_status'] = '✓ Trùng khớp'
                web_row['f5_type'] = 'exact'
            else:
                pct = f"{best[2]*100:.0f}%"
                ws[f'AB{row_num}'] = f'⚠ Gần khớp ({pct})'
                ws[f'AB{row_num}'].fill = light_orange
                stats['f5_partial'] += 1
                warnings.append(f'QT-BH: Trùng khớp một phần ({pct})')
                web_row['f5_status'] = f'⚠ Gần khớp ({pct})'
                web_row['f5_type'] = 'partial'
            if len(matches) > 1:
                ws[f'AB{row_num}'] = f'⚠ Nhiều KQ ({len(matches)})'
                ws[f'AB{row_num}'].fill = light_orange
                stats['f5_multiple'] += 1
                warnings.append(f'QT-BH: {len(matches)} kết quả khớp')
                web_row['f5_status'] = f'⚠ Nhiều KQ ({len(matches)})'
                web_row['f5_type'] = 'multiple'
                for m in matches:
                    multi_detail.append((rec['stt_goc'], ten_goc, 'QT-BH',
                                         m[0]['ten_quy_trinh'], m[0]['stt'], m[2]))
        else:
            ws[f'AB{row_num}'] = '✗ Chưa tìm thấy QT'
            ws[f'AB{row_num}'].fill = light_red
            stats['f5_not_found'] += 1
            warnings.append('QT-BH: Chưa tìm thấy quy trình')
            web_row['f5_status'] = '✗ Chưa tìm thấy'
            web_row['f5_type'] = 'not_found'

        # ---- Cảnh báo tổng hợp ----
        if warnings:
            ws[f'AC{row_num}'] = '\n'.join(warnings)
            ws[f'AC{row_num}'].fill = light_yellow
            if any('Chưa tìm thấy' in w for w in warnings):
                ws[f'AC{row_num}'].fill = light_red
            stats['need_review'] += 1
        else:
            ws[f'AC{row_num}'] = '✓ Đầy đủ'
            ws[f'AC{row_num}'].fill = light_green

        for col_idx in range(1, 30):
            cell = ws.cell(row=row_num, column=col_idx)
            if not cell.font or not cell.font.bold:
                cell.font = data_font
            cell.alignment = wrap_alignment
            cell.border = thin_border

        if warnings:
            detail_rows.append({
                'stt': rec['stt_goc'],
                'ten_goc': ten_goc,
                'warnings': warnings
            })

        web_row['warnings'] = warnings
        web_results.append(web_row)

        if (idx + 1) % 50 == 0 or idx + 1 == total:
            pct_done = (idx + 1) / total * 100
            log(f"    Đã xử lý {idx + 1}/{total} bản ghi ({pct_done:.0f}%)")
            if progress_callback:
                progress_callback(idx + 1, total)

    # Column widths
    col_widths = {
        'A': 6, 'B': 12, 'C': 40, 'D': 10, 'E': 10,
        'F': 12, 'G': 40, 'H': 12, 'I': 40,
        'J': 16, 'K': 14, 'L': 40, 'M': 10, 'N': 12,
        'O': 6, 'P': 14, 'Q': 18, 'R': 40, 'S': 30, 'T': 40, 'U': 30,
        'V': 6, 'W': 40,
        'X': 18, 'Y': 18, 'Z': 18, 'AA': 18, 'AB': 22, 'AC': 50,
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width
    ws.auto_filter.ref = f'A5:AC{5 + len(goc_records)}'

    # ================================================================
    # SHEET 2: THỐNG KÊ
    # ================================================================
    ws2 = wb.create_sheet("THONG KE")
    ws2.merge_cells('A1:F1')
    ws2['A1'] = 'BÁO CÁO THỐNG KÊ KẾT QUẢ ĐỐI CHIẾU'
    ws2['A1'].font = Font(name='Times New Roman', bold=True, size=14, color='2F5496')
    ws2['A1'].alignment = Alignment(horizontal='center')

    ws2['A3'] = 'Tổng số danh mục kỹ thuật trong file gốc:'
    ws2['A3'].font = data_font_bold
    ws2['D3'] = stats['total']
    ws2['D3'].font = Font(name='Times New Roman', bold=True, size=14, color='2F5496')

    stat_headers = ['Nguồn đối chiếu', 'Trùng khớp', 'Gần khớp\n(cần xác minh)',
                    'Nhiều kết quả', 'Chưa tìm thấy', 'Tỷ lệ khớp (%)']
    for i, h in enumerate(stat_headers):
        cell = ws2.cell(row=5, column=i+1, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_alignment
        cell.border = thin_border

    stat_data = [
        ('PL01 - TT23 (Mã & Tên KT)', stats['f1_found'], stats['f1_partial'],
         stats['f1_multiple'], stats['f1_not_found']),
        ('PL02 - TT23 (Mã liên kết & Tên KT)', stats['f2_found'], stats['f2_partial'],
         stats['f2_multiple'], stats['f2_not_found']),
        ('DMKT Tương đương (TT43/50/21)', stats['f3_found'], stats['f3_partial'],
         stats['f3_multiple'], stats['f3_not_found']),
        ('PL03 - DVKT (Phê duyệt)', stats['f4_found'], stats['f4_partial'],
         stats['f4_multiple'], stats['f4_not_found']),
        ('Tổng hợp QT đã ban hành', stats['f5_found'], stats['f5_partial'],
         stats['f5_multiple'], stats['f5_not_found']),
    ]

    for r_idx, (name, found, partial, multi, not_found) in enumerate(stat_data):
        row = 6 + r_idx
        ws2.cell(row=row, column=1, value=name).font = data_font_bold
        ws2.cell(row=row, column=2, value=found).font = data_font
        ws2.cell(row=row, column=3, value=partial).font = data_font
        ws2.cell(row=row, column=4, value=multi).font = data_font
        ws2.cell(row=row, column=5, value=not_found).font = data_font
        total_found = found + partial
        pct = f"{total_found / stats['total'] * 100:.1f}" if stats['total'] > 0 else "0"
        ws2.cell(row=row, column=6, value=f"{pct}%").font = data_font
        for c in range(1, 7):
            ws2.cell(row=row, column=c).border = thin_border
            ws2.cell(row=row, column=c).alignment = center_alignment
        if found > 0:
            ws2.cell(row=row, column=2).fill = light_green
        if not_found > 0:
            ws2.cell(row=row, column=5).fill = light_red
        if partial > 0:
            ws2.cell(row=row, column=3).fill = light_orange

    summary_row = 13
    ws2[f'A{summary_row}'] = 'TỔNG HỢP CẢNH BÁO'
    ws2[f'A{summary_row}'].font = Font(name='Times New Roman', bold=True, size=12, color='C00000')
    summary_items = [
        ('Số danh mục cần rà soát thủ công:', stats['need_review']),
        ('Số danh mục thiếu hồ sơ phê duyệt/giá/quy trình:', stats['missing_approval']),
        ('Số danh mục chưa tìm thấy tại PL01-TT23:', stats['f1_not_found']),
        ('Số danh mục chưa tìm thấy tại PL02-TT23:', stats['f2_not_found']),
        ('Số danh mục chưa tìm thấy tại DMKT Tương đương:', stats['f3_not_found']),
        ('Số danh mục chưa tìm thấy tại PL03-DVKT:', stats['f4_not_found']),
        ('Số danh mục chưa tìm thấy Quy trình đã BH:', stats['f5_not_found']),
    ]
    for i, (label, value) in enumerate(summary_items):
        ws2[f'A{summary_row + 1 + i}'] = label
        ws2[f'A{summary_row + 1 + i}'].font = data_font
        ws2[f'D{summary_row + 1 + i}'] = value
        ws2[f'D{summary_row + 1 + i}'].font = data_font_bold
        if value > 0:
            ws2[f'D{summary_row + 1 + i}'].fill = light_red
    ws2.column_dimensions['A'].width = 45
    for c in ['B', 'C', 'D', 'E', 'F']:
        ws2.column_dimensions[c].width = 18

    # ================================================================
    # SHEET 3: CẢNH BÁO CHI TIẾT
    # ================================================================
    ws3 = wb.create_sheet("CANH BAO CHI TIET")
    ws3.merge_cells('A1:D1')
    ws3['A1'] = 'DANH SÁCH CÁC DANH MỤC KỸ THUẬT CẦN RÀ SOÁT'
    ws3['A1'].font = Font(name='Times New Roman', bold=True, size=14, color='C00000')
    ws3['A1'].alignment = Alignment(horizontal='center')
    warn_headers = ['STT (Gốc)', 'Tên kỹ thuật (Gốc)', 'Nội dung cảnh báo', 'Ghi chú xử lý']
    for i, h in enumerate(warn_headers):
        cell = ws3.cell(row=3, column=i+1, value=h)
        cell.font = header_font
        cell.fill = PatternFill(start_color='C00000', end_color='C00000', fill_type='solid')
        cell.alignment = center_alignment
        cell.border = thin_border
    ws3.freeze_panes = 'A4'
    for r_idx, detail in enumerate(detail_rows):
        row = 4 + r_idx
        ws3.cell(row=row, column=1, value=detail['stt']).font = data_font
        ws3.cell(row=row, column=2, value=detail['ten_goc']).font = data_font
        ws3.cell(row=row, column=3, value='\n'.join(detail['warnings'])).font = data_font
        ws3.cell(row=row, column=4, value='').font = data_font
        for c in range(1, 5):
            ws3.cell(row=row, column=c).alignment = wrap_alignment
            ws3.cell(row=row, column=c).border = thin_border
        if any('Chưa tìm thấy' in w for w in detail['warnings']):
            for c in range(1, 5):
                ws3.cell(row=row, column=c).fill = light_red
        elif any('Thiếu' in w for w in detail['warnings']):
            for c in range(1, 5):
                ws3.cell(row=row, column=c).fill = light_orange
        else:
            for c in range(1, 5):
                ws3.cell(row=row, column=c).fill = light_yellow
    ws3.column_dimensions['A'].width = 10
    ws3.column_dimensions['B'].width = 50
    ws3.column_dimensions['C'].width = 60
    ws3.column_dimensions['D'].width = 30
    ws3.auto_filter.ref = f'A3:D{3 + len(detail_rows)}'

    # ================================================================
    # SHEET 4: NHIỀU KẾT QUẢ KHỚP
    # ================================================================
    ws4 = wb.create_sheet("NHIEU KET QUA KHOP")
    ws4.merge_cells('A1:F1')
    ws4['A1'] = 'DANH SÁCH CÁC DANH MỤC CÓ NHIỀU KẾT QUẢ ĐỐI CHIẾU'
    ws4['A1'].font = Font(name='Times New Roman', bold=True, size=14, color='BF8F00')
    ws4['A1'].alignment = Alignment(horizontal='center')
    multi_headers = ['STT (Gốc)', 'Tên kỹ thuật (Gốc)', 'File đối chiếu',
                     'Tên KT tìm thấy', 'Mã KT', 'Tỷ lệ khớp (%)']
    for i, h in enumerate(multi_headers):
        cell = ws4.cell(row=3, column=i+1, value=h)
        cell.font = header_font
        cell.fill = PatternFill(start_color='BF8F00', end_color='BF8F00', fill_type='solid')
        cell.alignment = center_alignment
        cell.border = thin_border
    ws4.freeze_panes = 'A4'
    for r_idx, (stt, ten_goc, file_label, ten_kt, ma_kt, score) in enumerate(multi_detail):
        row = 4 + r_idx
        ws4.cell(row=row, column=1, value=stt).font = data_font
        ws4.cell(row=row, column=2, value=ten_goc).font = data_font
        ws4.cell(row=row, column=3, value=file_label).font = data_font
        ws4.cell(row=row, column=4, value=ten_kt).font = data_font
        ws4.cell(row=row, column=5, value=ma_kt).font = data_font
        ws4.cell(row=row, column=6, value=f"{score*100:.1f}%").font = data_font
        for c in range(1, 7):
            ws4.cell(row=row, column=c).alignment = wrap_alignment
            ws4.cell(row=row, column=c).border = thin_border
    ws4.column_dimensions['A'].width = 10
    ws4.column_dimensions['B'].width = 50
    ws4.column_dimensions['C'].width = 15
    ws4.column_dimensions['D'].width = 50
    ws4.column_dimensions['E'].width = 18
    ws4.column_dimensions['F'].width = 15

    log("  Đang lưu file kết quả...")
    wb.save(output_path)
    log(f"  ✓ Đã lưu kết quả tại: {output_path}")
    return stats, web_results


# ============================================================================
# FLASK ROUTES
# ============================================================================

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')


@app.route('/api/upload', methods=['POST'])
def upload_files():
    """Handle file upload and start processing."""
    session_id = str(uuid.uuid4())

    file_keys = ['file_goc', 'file_1', 'file_2', 'file_3', 'file_4', 'file_5']
    file_labels = [
        'File gốc (Phụ lục XII)',
        'File 1 - PL01 TT23',
        'File 2 - PL02 TT23',
        'File 3 - DMKT Tương đương',
        'File 4 - PL03 DVKT',
        'File 5 - Tổng hợp QT',
    ]

    saved_paths = {}
    for key, label in zip(file_keys, file_labels):
        if key not in request.files:
            return jsonify({'error': f'Thiếu {label}'}), 400
        f = request.files[key]
        if f.filename == '':
            return jsonify({'error': f'Chưa chọn {label}'}), 400

        session_dir = os.path.join(UPLOAD_DIR, session_id)
        os.makedirs(session_dir, exist_ok=True)
        save_path = os.path.join(session_dir, f'{key}_{f.filename}')
        f.save(save_path)
        saved_paths[key] = save_path

    # Initialize session
    sessions[session_id] = {
        'status': 'processing',
        'progress': 0,
        'total': 0,
        'logs': [],
        'stats': None,
        'web_results': None,
        'result_file': None,
        'error': None,
    }

    # Start background processing
    thread = threading.Thread(
        target=process_files,
        args=(session_id, saved_paths),
        daemon=True
    )
    thread.start()

    return jsonify({'session_id': session_id})


def process_files(session_id, paths):
    """Background processing worker."""
    session = sessions[session_id]

    def log_cb(msg):
        session['logs'].append(msg)

    def progress_cb(current, total):
        session['progress'] = current
        session['total'] = total

    try:
        log_cb("Bước 1: Đọc dữ liệu...")

        log_cb("  Đọc file gốc (Phụ lục XII)...")
        goc_records = read_file_goc(paths['file_goc'])
        log_cb(f"  ✓ File gốc: {len(goc_records)} danh mục kỹ thuật")

        log_cb("  Đọc File 1 (PL01 TT23)...")
        data_f1 = read_file_1(paths['file_1'])
        log_cb(f"  ✓ File 1: {len(data_f1)} bản ghi")

        log_cb("  Đọc File 2 (PL02 TT23)...")
        data_f2 = read_file_2(paths['file_2'])
        log_cb(f"  ✓ File 2: {len(data_f2)} bản ghi")

        log_cb("  Đọc File 3 (DMKT Tương đương)...")
        data_f3 = read_file_3(paths['file_3'])
        log_cb(f"  ✓ File 3: {len(data_f3)} bản ghi")

        log_cb("  Đọc File 4 (PL03 DVKT)...")
        data_f4 = read_file_4(paths['file_4'])
        log_cb(f"  ✓ File 4: {len(data_f4)} bản ghi")

        log_cb("  Đọc File 5 (Tổng hợp QT)...")
        data_f5 = read_file_5(paths['file_5'])
        log_cb(f"  ✓ File 5: {len(data_f5)} bản ghi")

        log_cb("")
        log_cb("Bước 2: Đối chiếu và tổng hợp...")

        output_filename = f'KET_QUA_DOI_CHIEU_{session_id[:8]}.xlsx'
        output_path = os.path.join(RESULT_DIR, output_filename)

        stats, web_results = create_result_workbook(
            goc_records, data_f1, data_f2, data_f3, data_f4, data_f5,
            output_path,
            log_callback=log_cb,
            progress_callback=progress_cb
        )

        session['stats'] = stats
        session['web_results'] = web_results
        session['result_file'] = output_filename
        session['status'] = 'completed'
        log_cb("")
        log_cb("✓ HOÀN THÀNH!")

    except Exception as e:
        session['status'] = 'error'
        session['error'] = str(e)
        log_cb(f"LỖI: {str(e)}")
        import traceback
        log_cb(traceback.format_exc())


@app.route('/api/status/<session_id>')
def get_status(session_id):
    """Get processing status."""
    if session_id not in sessions:
        return jsonify({'error': 'Session không tồn tại'}), 404

    session = sessions[session_id]
    return jsonify({
        'status': session['status'],
        'progress': session['progress'],
        'total': session['total'],
        'logs': session['logs'],
        'error': session['error'],
    })


@app.route('/api/results/<session_id>')
def get_results(session_id):
    """Get processing results."""
    if session_id not in sessions:
        return jsonify({'error': 'Session không tồn tại'}), 404

    session = sessions[session_id]
    if session['status'] != 'completed':
        return jsonify({'error': 'Chưa xử lý xong'}), 400

    return jsonify({
        'stats': session['stats'],
        'results': session['web_results'],
        'result_file': session['result_file'],
    })


@app.route('/api/download/<filename>')
def download_file(filename):
    """Download result file."""
    filepath = os.path.join(RESULT_DIR, filename)
    if not os.path.exists(filepath):
        return jsonify({'error': 'File không tồn tại'}), 404
    return send_file(filepath, as_attachment=True, download_name='KET_QUA_DOI_CHIEU_DANH_MUC_KY_THUAT.xlsx')


@app.route('/api/stream/<session_id>')
def stream_progress(session_id):
    """SSE endpoint for real-time progress."""
    def generate():
        last_log_idx = 0
        while True:
            if session_id not in sessions:
                yield f"data: {json.dumps({'type': 'error', 'message': 'Session not found'})}\n\n"
                break

            session = sessions[session_id]
            new_logs = session['logs'][last_log_idx:]
            last_log_idx = len(session['logs'])

            if new_logs:
                for log_msg in new_logs:
                    yield f"data: {json.dumps({'type': 'log', 'message': log_msg})}\n\n"

            if session['total'] > 0:
                yield f"data: {json.dumps({'type': 'progress', 'current': session['progress'], 'total': session['total']})}\n\n"

            if session['status'] == 'completed':
                yield f"data: {json.dumps({'type': 'completed', 'result_file': session['result_file']})}\n\n"
                break
            elif session['status'] == 'error':
                yield f"data: {json.dumps({'type': 'error', 'message': session.get('error', 'Unknown error')})}\n\n"
                break

            time.sleep(0.5)

    return Response(generate(), mimetype='text/event-stream',
                    headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'})


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)
